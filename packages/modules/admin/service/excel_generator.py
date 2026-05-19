"""Excel report generator for tenant exports.

Generates client-friendly Excel files for non-technical users (accountants, admins).
Uses openpyxl library with Spanish headers for client-facing reports.
"""

from io import BytesIO
from typing import Any

from sqlalchemy.orm import Session

# Check openpyxl availability
try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    Workbook = None


class ExcelGeneratorError(Exception):
    """Raised when Excel generation fails."""
    pass


class ExcelGenerator:
    """Generate Excel reports for tenant exports.

    Produces formatted Excel files with:
    - Blue header rows with white bold text
    - Auto-sized columns
    - Currency formatting for amounts
    - Spanish headers for client-facing reports
    """

    HEADER_FILL = PatternFill(
        start_color="4472C4",
        end_color="4472C4",
        fill_type="solid"
    )
    HEADER_FONT = Font(bold=True, color="FFFFFF")
    CURRENCY_FORMAT = '$#,##0.00'

    def __init__(self, db: Session):
        """Initialize Excel generator.

        Args:
            db: SQLAlchemy session for database queries

        Raises:
            ExcelGeneratorError: If openpyxl is not available
        """
        if not OPENPYXL_AVAILABLE:
            raise ExcelGeneratorError(
                "openpyxl is required for Excel generation. "
                "Install with: pip install openpyxl"
            )
        self.db = db

    def _create_workbook(self) -> Workbook:
        """Create a new openpyxl workbook."""
        return Workbook()

    def _auto_adjust_columns(self, ws):
        """Auto-size columns based on content.

        Args:
            ws: Worksheet to adjust
        """
        for column_cells in ws.columns:
            max_length = 0
            column = column_cells[0].column_letter
            for cell in column_cells:
                try:
                    if cell.value:
                        cell_length = len(str(cell.value))
                        if cell_length > max_length:
                            max_length = cell_length
                except Exception:
                    pass
            # Add padding for readability
            adjusted_width = min(max_length + 2, 50)
            ws.column_dimensions[column].width = adjusted_width

    def _add_header(self, ws, headers: list[str], row: int = 1):
        """Add styled header row.

        Blue fill, white bold text, centered alignment.

        Args:
            ws: Worksheet to add header to
            headers: List of header labels
            row: Row number for header (default 1)
        """
        for col, header in enumerate(headers, start=1):
            cell = ws.cell(row=row, column=col, value=header)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def generate_expenses_excel(
        self,
        company_id: int,
        expenses: list[Any],
        *,
        include_cfdi: bool = True
    ) -> bytes:
        """Generate Excel report for expenses.

        Args:
            company_id: Company ID for context
            expenses: List of expense objects with attributes
            include_cfdi: Whether to include CFDI columns

        Returns:
            Excel file as bytes
        """
        wb = self._create_workbook()
        ws = wb.active
        ws.title = "Gastos"

        # Headers in Spanish
        headers = [
            "ID",
            "Fecha",
            "Descripción",
            "Monto",
            "Moneda",
            "Categoría",
            "Proveedor",
            "RFC Proveedor",
            "Estado",
            "Tipo de Asentamiento",
        ]
        if include_cfdi:
            headers.append("UUID CFDI")
        headers.extend(["Notas", "Creado"])

        self._add_header(ws, headers)

        # Data rows
        for row_idx, expense in enumerate(expenses, start=2):
            ws.cell(row=row_idx, column=1, value=getattr(expense, 'id', ''))
            ws.cell(row=row_idx, column=2, value=getattr(expense, 'expense_date', ''))
            ws.cell(row=row_idx, column=3, value=getattr(expense, 'description', ''))
            # Amount with currency format
            amount_cell = ws.cell(row=row_idx, column=4, value=getattr(expense, 'amount', 0))
            amount_cell.number_format = self.CURRENCY_FORMAT
            ws.cell(row=row_idx, column=5, value=getattr(expense, 'currency', 'MXN'))
            ws.cell(row=row_idx, column=6, value=getattr(expense, 'category_name', ''))
            ws.cell(row=row_idx, column=7, value=getattr(expense, 'vendor_name', ''))
            ws.cell(row=row_idx, column=8, value=getattr(expense, 'vendor_rfc', ''))
            ws.cell(row=row_idx, column=9, value=getattr(expense, 'status', ''))
            ws.cell(row=row_idx, column=10, value=getattr(expense, 'settlement_type', ''))

            col = 11
            if include_cfdi:
                ws.cell(row=row_idx, column=col, value=getattr(expense, 'cfdi_uuid', ''))
                col += 1
            ws.cell(row=row_idx, column=col, value=getattr(expense, 'notes', ''))
            ws.cell(row=row_idx, column=col + 1, value=getattr(expense, 'created_at', ''))

        self._auto_adjust_columns(ws)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    def generate_categories_excel(
        self,
        company_id: int,
        categories: list[Any]
    ) -> bytes:
        """Generate Excel report for accounting categories.

        Args:
            company_id: Company ID for context
            categories: List of category objects

        Returns:
            Excel file as bytes
        """
        wb = self._create_workbook()
        ws = wb.active
        ws.title = "Categorías"

        headers = ["Código", "Nombre", "Descripción", "Tipo", "Cuenta Contable"]
        self._add_header(ws, headers)

        for row_idx, cat in enumerate(categories, start=2):
            ws.cell(row=row_idx, column=1, value=getattr(cat, 'code', ''))
            ws.cell(row=row_idx, column=2, value=getattr(cat, 'name', ''))
            ws.cell(row=row_idx, column=3, value=getattr(cat, 'description', ''))
            ws.cell(row=row_idx, column=4, value=getattr(cat, 'type', ''))
            ws.cell(row=row_idx, column=5, value=getattr(cat, 'account_code', ''))

        self._auto_adjust_columns(ws)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    def generate_vendors_excel(
        self,
        company_id: int,
        vendors: list[dict]
    ) -> bytes:
        """Generate Excel report for vendors/suppliers.

        Args:
            company_id: Company ID for context
            vendors: List of vendor dictionaries with aggregated data

        Returns:
            Excel file as bytes
        """
        wb = self._create_workbook()
        ws = wb.active
        ws.title = "Proveedores"

        headers = ["RFC", "Nombre", "Total Gastos", "Monto Total", "Último Gasto"]
        self._add_header(ws, headers)

        for row_idx, vendor in enumerate(vendors, start=2):
            ws.cell(row=row_idx, column=1, value=vendor.get('rfc', ''))
            ws.cell(row=row_idx, column=2, value=vendor.get('name', ''))
            ws.cell(row=row_idx, column=3, value=vendor.get('expense_count', 0))
            # Total amount with currency format
            amount_cell = ws.cell(row=row_idx, column=4, value=vendor.get('total_amount', 0))
            amount_cell.number_format = self.CURRENCY_FORMAT
            ws.cell(row=row_idx, column=5, value=vendor.get('last_expense_date', ''))

        self._auto_adjust_columns(ws)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    def generate_audit_excel(self, audit_data: list[dict]) -> bytes:
        """Generate Excel report for audit events.

        Args:
            audit_data: List of audit event dictionaries

        Returns:
            Excel file as bytes
        """
        wb = self._create_workbook()
        ws = wb.active
        ws.title = "Auditoría"

        headers = [
            "ID",
            "Fecha/Hora",
            "Tipo de Evento",
            "Tipo de Entidad",
            "ID de Entidad",
            "ID de Usuario",
            "IP",
            "Datos Antes",
            "Datos Después"
        ]
        self._add_header(ws, headers)

        for row_idx, event in enumerate(audit_data, start=2):
            ws.cell(row=row_idx, column=1, value=event.get('id', ''))
            ws.cell(row=row_idx, column=2, value=event.get('timestamp', ''))
            ws.cell(row=row_idx, column=3, value=event.get('event_type', ''))
            ws.cell(row=row_idx, column=4, value=event.get('entity_type', ''))
            ws.cell(row=row_idx, column=5, value=event.get('entity_id', ''))
            ws.cell(row=row_idx, column=6, value=event.get('user_id', ''))
            ws.cell(row=row_idx, column=7, value=event.get('ip_address', ''))
            ws.cell(row=row_idx, column=8, value=event.get('before_data', ''))
            ws.cell(row=row_idx, column=9, value=event.get('after_data', ''))

        self._auto_adjust_columns(ws)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()

    def generate_accounting_excel(
        self,
        company_id: int,
        outputs: list[Any]
    ) -> bytes:
        """Generate Excel report for accounting outputs (Polizas).

        Args:
            company_id: Company ID for context
            outputs: List of accounting output objects

        Returns:
            Excel file as bytes
        """
        wb = self._create_workbook()
        ws = wb.active
        ws.title = "Contabilidad"

        headers = [
            "ID",
            "Fecha",
            "Tipo",
            "Descripción",
            "Total Débitos",
            "Total Créditos",
            "Estado"
        ]
        self._add_header(ws, headers)

        for row_idx, output in enumerate(outputs, start=2):
            ws.cell(row=row_idx, column=1, value=getattr(output, 'id', ''))
            ws.cell(row=row_idx, column=2, value=getattr(output, 'date', ''))
            ws.cell(row=row_idx, column=3, value=getattr(output, 'type', ''))
            ws.cell(row=row_idx, column=4, value=getattr(output, 'description', ''))
            # Debits with currency format
            debit_cell = ws.cell(row=row_idx, column=5, value=getattr(output, 'total_debits', 0))
            debit_cell.number_format = self.CURRENCY_FORMAT
            # Credits with currency format
            credit_cell = ws.cell(row=row_idx, column=6, value=getattr(output, 'total_credits', 0))
            credit_cell.number_format = self.CURRENCY_FORMAT
            ws.cell(row=row_idx, column=7, value=getattr(output, 'status', ''))

        self._auto_adjust_columns(ws)

        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return output.getvalue()