"""Tests for Excel report generator."""

import pytest
from io import BytesIO
from unittest.mock import MagicMock, patch

from packages.modules.admin.service.excel_generator import (
    ExcelGenerator,
    ExcelGeneratorError,
    OPENPYXL_AVAILABLE,
)


@pytest.fixture
def mock_db():
    """Create mock database session."""
    return MagicMock()


@pytest.fixture
def excel_generator(mock_db):
    """Create ExcelGenerator instance."""
    if not OPENPYXL_AVAILABLE:
        pytest.skip("openpyxl not installed")
    return ExcelGenerator(mock_db)


class MockExpense:
    """Mock expense object for testing."""
    def __init__(self, **kwargs):
        self.id = kwargs.get('id', 1)
        self.expense_date = kwargs.get('expense_date', '2024-01-15')
        self.description = kwargs.get('description', 'Office supplies')
        self.amount = kwargs.get('amount', 1500.50)
        self.currency = kwargs.get('currency', 'MXN')
        self.category_name = kwargs.get('category_name', 'Office Expenses')
        self.vendor_name = kwargs.get('vendor_name', 'Office Depot')
        self.vendor_rfc = kwargs.get('vendor_rfc', 'ODE123456ABC')
        self.status = kwargs.get('status', 'approved')
        self.settlement_type = kwargs.get('settlement_type', 'reimbursement')
        self.cfdi_uuid = kwargs.get('cfdi_uuid', 'abc-123-def')
        self.notes = kwargs.get('notes', 'Monthly supplies')
        self.created_at = kwargs.get('created_at', '2024-01-10')


class MockCategory:
    """Mock category object for testing."""
    def __init__(self, **kwargs):
        self.code = kwargs.get('code', 'CAT001')
        self.name = kwargs.get('name', 'Office Expenses')
        self.description = kwargs.get('description', 'Office related expenses')
        self.type = kwargs.get('type', 'expense')
        self.account_code = kwargs.get('account_code', '600-01')


class MockAccountingOutput:
    """Mock accounting output object for testing."""
    def __init__(self, **kwargs):
        self.id = kwargs.get('id', 1)
        self.date = kwargs.get('date', '2024-01-15')
        self.type = kwargs.get('type', 'egreso')
        self.description = kwargs.get('description', 'Reembolso gastos oficina')
        self.total_debits = kwargs.get('total_debits', 1500.50)
        self.total_credits = kwargs.get('total_credits', 1500.50)
        self.status = kwargs.get('status', 'posted')


class TestExcelGeneratorInit:
    """Tests for ExcelGenerator initialization."""

    def test_init_success(self, mock_db):
        """Test successful initialization when openpyxl is available."""
        if not OPENPYXL_AVAILABLE:
            pytest.skip("openpyxl not installed")
        generator = ExcelGenerator(mock_db)
        assert generator.db == mock_db

    def test_init_without_openpyxl(self, mock_db):
        """Test initialization fails when openpyxl is not available."""
        with patch('packages.modules.admin.service.excel_generator.OPENPYXL_AVAILABLE', False):
            with pytest.raises(ExcelGeneratorError) as exc_info:
                ExcelGenerator(mock_db)
            assert "openpyxl is required" in str(exc_info.value)


class TestExpensesExcel:
    """Tests for expenses Excel generation."""

    def test_empty_expenses_generates_valid_excel(self, excel_generator):
        """Test that empty data generates valid Excel file."""
        result = excel_generator.generate_expenses_excel(
            company_id=1,
            expenses=[]
        )
        # Excel files start with PK (ZIP signature for .xlsx)
        assert result.startswith(b"PK")
        assert len(result) > 0

    def test_non_empty_expenses_generates_valid_excel(self, excel_generator):
        """Test that non-empty data generates valid Excel file."""
        expenses = [
            MockExpense(id=1, description="Test expense 1", amount=100.00),
            MockExpense(id=2, description="Test expense 2", amount=200.00),
        ]
        result = excel_generator.generate_expenses_excel(
            company_id=1,
            expenses=expenses
        )
        assert result.startswith(b"PK")
        assert len(result) > 100  # Should have meaningful content

    def test_expenses_excel_with_cfdi(self, excel_generator):
        """Test expenses Excel includes CFDI column when enabled."""
        expenses = [
            MockExpense(id=1, cfdi_uuid='test-uuid-123')
        ]
        result = excel_generator.generate_expenses_excel(
            company_id=1,
            expenses=expenses,
            include_cfdi=True
        )
        assert result.startswith(b"PK")

    def test_expenses_excel_without_cfdi(self, excel_generator):
        """Test expenses Excel excludes CFDI column when disabled."""
        expenses = [
            MockExpense(id=1, cfdi_uuid='test-uuid-123')
        ]
        result = excel_generator.generate_expenses_excel(
            company_id=1,
            expenses=expenses,
            include_cfdi=False
        )
        assert result.startswith(b"PK")


class TestCategoriesExcel:
    """Tests for categories Excel generation."""

    def test_empty_categories_generates_valid_excel(self, excel_generator):
        """Test that empty data generates valid Excel file."""
        result = excel_generator.generate_categories_excel(
            company_id=1,
            categories=[]
        )
        assert result.startswith(b"PK")
        assert len(result) > 0

    def test_non_empty_categories_generates_valid_excel(self, excel_generator):
        """Test that non-empty data generates valid Excel file."""
        categories = [
            MockCategory(code='CAT001', name='Category 1'),
            MockCategory(code='CAT002', name='Category 2'),
        ]
        result = excel_generator.generate_categories_excel(
            company_id=1,
            categories=categories
        )
        assert result.startswith(b"PK")
        assert len(result) > 100


class TestVendorsExcel:
    """Tests for vendors Excel generation."""

    def test_empty_vendors_generates_valid_excel(self, excel_generator):
        """Test that empty data generates valid Excel file."""
        result = excel_generator.generate_vendors_excel(
            company_id=1,
            vendors=[]
        )
        assert result.startswith(b"PK")
        assert len(result) > 0

    def test_non_empty_vendors_generates_valid_excel(self, excel_generator):
        """Test that non-empty data generates valid Excel file."""
        vendors = [
            {
                'rfc': 'ABC123456XYZ',
                'name': 'Vendor 1',
                'expense_count': 10,
                'total_amount': 5000.00,
                'last_expense_date': '2024-01-15'
            },
            {
                'rfc': 'DEF789012UVW',
                'name': 'Vendor 2',
                'expense_count': 5,
                'total_amount': 2500.00,
                'last_expense_date': '2024-01-10'
            },
        ]
        result = excel_generator.generate_vendors_excel(
            company_id=1,
            vendors=vendors
        )
        assert result.startswith(b"PK")
        assert len(result) > 100


class TestAuditExcel:
    """Tests for audit Excel generation."""

    def test_empty_audit_generates_valid_excel(self, excel_generator):
        """Test that empty data generates valid Excel file."""
        result = excel_generator.generate_audit_excel(audit_data=[])
        assert result.startswith(b"PK")
        assert len(result) > 0

    def test_non_empty_audit_generates_valid_excel(self, excel_generator):
        """Test that non-empty data generates valid Excel file."""
        audit_data = [
            {
                'id': 1,
                'timestamp': '2024-01-15 10:30:00',
                'event_type': 'create',
                'entity_type': 'expense',
                'entity_id': 100,
                'user_id': 5,
                'ip_address': '192.168.1.1',
                'before_data': None,
                'after_data': '{"status": "draft"}'
            },
        ]
        result = excel_generator.generate_audit_excel(audit_data=audit_data)
        assert result.startswith(b"PK")
        assert len(result) > 100


class TestAccountingExcel:
    """Tests for accounting Excel generation."""

    def test_empty_accounting_generates_valid_excel(self, excel_generator):
        """Test that empty data generates valid Excel file."""
        result = excel_generator.generate_accounting_excel(
            company_id=1,
            outputs=[]
        )
        assert result.startswith(b"PK")
        assert len(result) > 0

    def test_non_empty_accounting_generates_valid_excel(self, excel_generator):
        """Test that non-empty data generates valid Excel file."""
        outputs = [
            MockAccountingOutput(id=1, description='Poliza 1'),
            MockAccountingOutput(id=2, description='Poliza 2'),
        ]
        result = excel_generator.generate_accounting_excel(
            company_id=1,
            outputs=outputs
        )
        assert result.startswith(b"PK")
        assert len(result) > 100


class TestExcelStructure:
    """Tests for Excel file structure and formatting."""

    def test_excel_can_be_opened(self, excel_generator):
        """Test that generated Excel can be parsed by openpyxl."""
        from openpyxl import load_workbook

        expenses = [MockExpense(id=1, description="Test")]
        result = excel_generator.generate_expenses_excel(
            company_id=1,
            expenses=expenses
        )

        # Verify it can be loaded
        wb = load_workbook(BytesIO(result))
        ws = wb.active
        assert ws.title == "Gastos"
        assert ws.max_row >= 1  # At least header

    def test_header_styling_applied(self, excel_generator):
        """Test that header row has proper styling."""
        from openpyxl import load_workbook

        result = excel_generator.generate_categories_excel(
            company_id=1,
            categories=[MockCategory()]
        )

        wb = load_workbook(BytesIO(result))
        ws = wb.active

        # Check first cell styling
        header_cell = ws.cell(row=1, column=1)
        assert header_cell.fill.start_color.rgb == "004472C4"
        assert header_cell.font.bold is True

    def test_spanish_headers_present(self, excel_generator):
        """Test that Spanish headers are used in reports."""
        from openpyxl import load_workbook

        result = excel_generator.generate_expenses_excel(
            company_id=1,
            expenses=[MockExpense()]
        )

        wb = load_workbook(BytesIO(result))
        ws = wb.active

        headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]

        # Check for Spanish headers
        assert "Descripción" in headers
        assert "Monto" in headers
        assert "Categoría" in headers
        assert "Proveedor" in headers

    def test_currency_formatting(self, excel_generator):
        """Test that currency cells have proper number format."""
        from openpyxl import load_workbook

        expenses = [MockExpense(id=1, amount=1234.56)]
        result = excel_generator.generate_expenses_excel(
            company_id=1,
            expenses=expenses
        )

        wb = load_workbook(BytesIO(result))
        ws = wb.active

        # Find the amount column (column 4)
        amount_cell = ws.cell(row=2, column=4)
        assert amount_cell.number_format == '$#,##0.00'